# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Angshuman Nandy
from __future__ import annotations

import asyncio
import logging
from functools import partial

from app.pipeline.base import BaseTask
from app.pipeline.context import Document, PipelineContext
from app.pipeline.registry import register

logger = logging.getLogger(__name__)


def _extract_text(file_path: str) -> str:
    """Open a DOCX file and return all paragraph text joined by newlines."""
    import docx  # python-docx

    document = docx.Document(file_path)
    paragraphs = [para.text for para in document.paragraphs]
    return "\n".join(paragraphs)


@register("ingestion", "docx")
class DOCXIngestionTask(BaseTask):
    """Ingest a DOCX file as a single Document."""

    async def run(self, context: PipelineContext) -> PipelineContext:
        file_path: str = context.metadata["file_path"]
        filename: str = context.metadata.get("filename", file_path)

        loop = asyncio.get_event_loop()
        text = await loop.run_in_executor(None, partial(_extract_text, file_path))

        doc = Document(
            content=text,
            metadata={
                "source": filename,
                "file_type": "docx",
            },
        )
        context.documents.append(doc)

        logger.debug("DOCXIngestionTask: extracted text from %r", filename)
        return context

    async def validate_config(self, config: dict) -> None:
        pass
